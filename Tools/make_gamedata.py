"""
DungeonColony 게임 데이터 Excel 생성기
실행: python Tools/make_gamedata.py
출력: Tools/GameData.xlsx
"""

from openpyxl import Workbook
from openpyxl.styles import (PatternFill, Font, Alignment, Border, Side,
                              GradientFill)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.formatting.rule import Rule
import os

OUTPUT_PATH = os.path.join(os.path.dirname(__file__), "GameData.xlsx")

# ── 팔레트 ─────────────────────────────────────────────────────────────────
C = {
    "header_bg":   "1F2D3D",
    "header_fg":   "FFFFFF",
    "section_ore": "D6E4F0",   # 광석 — 연파랑
    "section_gem": "F9E4D4",   # 보석 — 연주황
    "section_stone":"E8F5E9",  # 석재/목재 — 연초록
    "section_ingot":"FFF9C4",  # 주괴 — 연노랑
    "section_magic":"EDE7F6",  # 마법석 — 연보라
    "section_crystal":"FCE4EC",# 결정 — 연분홍
    "section_smelt":"FFF3E0",  # 제련 레시피 — 연주황
    "section_stone_r":"F1F8E9",# 석재가공 레시피
    "section_wood_r": "E8F5E9",# 목공 레시피
    "section_gem_r":  "F3E5F5",# 보석가공 레시피
    "section_weapon": "FFEBEE",# 무기제작
    "section_armor":  "E3F2FD",# 방어구제작
    "stripe":      "F5F5F5",
    "white":       "FFFFFF",
    "border":      "BDBDBD",
    "note":        "9E9E9E",
}

def fill(hex_color):
    # openpyxl requires 8-char ARGB; prepend FF if 6-char RGB given
    c = hex_color if len(hex_color) == 8 else "FF" + hex_color
    return PatternFill("solid", fgColor=c)

def hfont(bold=True, color="FFFFFF", size=10):
    return Font(bold=bold, color=color, name="맑은 고딕", size=size)

def cfont(bold=False, color="212121", size=10):
    return Font(bold=bold, color=color, name="맑은 고딕", size=size)

def border():
    s = Side(style="thin", color=C["border"])
    return Border(left=s, right=s, top=s, bottom=s)

def center():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)

def left():
    return Alignment(horizontal="left", vertical="center", wrap_text=True)

def apply_header(ws, headers, row=1):
    """헤더 행 스타일 적용"""
    for col, (label, width) in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=label)
        c.fill     = fill(C["header_bg"])
        c.font     = hfont()
        c.alignment= center()
        c.border   = border()
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[row].height = 20

def style_row(ws, row_num, n_cols, bg_color, bold=False):
    for col in range(1, n_cols + 1):
        c = ws.cell(row=row_num, column=col)
        c.fill      = fill(bg_color)
        c.font      = cfont(bold=bold)
        c.alignment = left() if col > 2 else center()
        c.border    = border()

def section_label(ws, row_num, n_cols, label, bg):
    ws.merge_cells(start_row=row_num, start_column=1,
                   end_row=row_num,   end_column=n_cols)
    c = ws.cell(row=row_num, column=1, value=f"  {label}")
    c.fill      = fill(bg)
    c.font      = cfont(bold=True, color="424242", size=9)
    c.alignment = left()
    c.border    = border()
    ws.row_dimensions[row_num].height = 16


# ══════════════════════════════════════════════════════════════════════════════
# 시트 1 : Items
# ══════════════════════════════════════════════════════════════════════════════
ITEM_HEADERS = [
    ("itemName",      22),
    ("type",          12),
    ("resourceType",  18),
    ("weight",         8),
    ("maxStack",       9),
    ("width",          7),
    ("height",         7),
    ("description",   40),
    ("iconPath",      25),
]

ITEM_TYPES = ["resource", "weapon", "armor", "seed", ""]
RESOURCE_TYPES = [
    "Stone","CutStone",
    "CopperOre","TinOre","IronOre","Coal","SilverOre","GoldOre",
    "MagicStone","Etherstone","Voidstone",
    "Ruby","Sapphire","Topaz","Emerald",
    "Bronze","Steel","SilverIngot","GoldIngot",
    "Wood","Lumber",
]

ITEMS = [
    # (itemName, type, resourceType, weight, maxStack, w, h, description, iconPath, bg_color)
    # ── 석재 / 목재 ──────────────────────────────────────────
    ("돌",          "resource","Stone",    1.0, 999,1,1,"채굴로 얻는 기본 석재","",         C["section_stone"]),
    ("다듬은 돌",   "resource","CutStone", 1.5, 999,1,1,"석재 가공대에서 다듬은 돌","",    C["section_stone"]),
    ("목재",        "resource","Wood",     0.8, 999,1,1,"벌목으로 얻는 나무","",            C["section_stone"]),
    ("목재 판자",   "resource","Lumber",   0.6, 999,1,1,"제재소에서 가공한 판자","",        C["section_stone"]),
    # ── 광석 ─────────────────────────────────────────────────
    ("구리 광석",   "resource","CopperOre", 1.0,999,1,1,"","", C["section_ore"]),
    ("주석 광석",   "resource","TinOre",    1.0,999,1,1,"","", C["section_ore"]),
    ("철 광석",     "resource","IronOre",   1.2,999,1,1,"","", C["section_ore"]),
    ("석탄",        "resource","Coal",      0.8,999,1,1,"용광로 연료로도 사용","", C["section_ore"]),
    ("은 광석",     "resource","SilverOre", 1.5,999,1,1,"","", C["section_ore"]),
    ("금 광석",     "resource","GoldOre",   2.0,999,1,1,"","", C["section_ore"]),
    # ── 주괴 ─────────────────────────────────────────────────
    ("청동 주괴",   "resource","Bronze",      2.0,999,1,1,"구리+주석 제련","", C["section_ingot"]),
    ("강철 주괴",   "resource","Steel",       2.5,999,1,1,"철+석탄 제련","",   C["section_ingot"]),
    ("은 주괴",     "resource","SilverIngot", 2.0,999,1,1,"","",               C["section_ingot"]),
    ("금 주괴",     "resource","GoldIngot",   2.5,999,1,1,"","",               C["section_ingot"]),
    # ── 마법석 ───────────────────────────────────────────────
    ("마력석",      "resource","MagicStone", 0.5,999,1,1,"원소 계열 마법 촉매","", C["section_magic"]),
    ("에테르석",    "resource","Etherstone", 0.5,999,1,1,"신성 계열 마법 촉매","",C["section_magic"]),
    ("공허석",      "resource","Voidstone",  0.5,999,1,1,"어둠 계열 마법 촉매","",C["section_magic"]),
    # ── 보석 ─────────────────────────────────────────────────
    ("루비",        "resource","Ruby",     0.3,999,1,1,"화염 속성 보석","", C["section_gem"]),
    ("사파이어",    "resource","Sapphire", 0.3,999,1,1,"냉기 속성 보석","", C["section_gem"]),
    ("토파즈",      "resource","Topaz",    0.3,999,1,1,"번개 속성 보석","", C["section_gem"]),
    ("에메랄드",    "resource","Emerald",  0.3,999,1,1,"독 속성 보석","",   C["section_gem"]),
    # ── 마법 결정 ────────────────────────────────────────────
    ("원소 화염 결정",  "","",0.2,1,1,1,"루비+마력석","",  C["section_crystal"]),
    ("원소 냉기 결정",  "","",0.2,1,1,1,"사파이어+마력석","",C["section_crystal"]),
    ("원소 번개 결정",  "","",0.2,1,1,1,"토파즈+마력석","", C["section_crystal"]),
    ("원소 독 결정",    "","",0.2,1,1,1,"에메랄드+마력석","",C["section_crystal"]),
    ("신성 화염 결정",  "","",0.2,1,1,1,"루비+에테르석","", C["section_crystal"]),
    ("신성 냉기 결정",  "","",0.2,1,1,1,"사파이어+에테르석","",C["section_crystal"]),
    ("신성 번개 결정",  "","",0.2,1,1,1,"토파즈+에테르석","",C["section_crystal"]),
    ("신성 독 결정",    "","",0.2,1,1,1,"에메랄드+에테르석","",C["section_crystal"]),
    ("암흑 화염 결정",  "","",0.2,1,1,1,"루비+공허석","",   C["section_crystal"]),
    ("암흑 냉기 결정",  "","",0.2,1,1,1,"사파이어+공허석","",C["section_crystal"]),
    ("암흑 번개 결정",  "","",0.2,1,1,1,"토파즈+공허석","", C["section_crystal"]),
    ("암흑 독 결정",    "","",0.2,1,1,1,"에메랄드+공허석","",C["section_crystal"]),
]


def make_items_sheet(wb):
    ws = wb.active
    ws.title = "Items"
    ws.freeze_panes = "A2"

    apply_header(ws, ITEM_HEADERS)

    # 드롭다운 — type
    dv_type = DataValidation(
        type="list",
        formula1='"' + ",".join(ITEM_TYPES) + '"',
        showDropDown=False,
        error="목록에서 선택하세요",
        errorTitle="입력 오류",
    )
    ws.add_data_validation(dv_type)

    # 드롭다운 — resourceType
    dv_rt = DataValidation(
        type="list",
        formula1='"' + ",".join(RESOURCE_TYPES) + '"',
        showDropDown=False,
    )
    ws.add_data_validation(dv_rt)

    SECTIONS = [
        ("▼ 석재 / 목재", C["section_stone"],  4),
        ("▼ 광석",         C["section_ore"],    6),
        ("▼ 주괴",         C["section_ingot"],  4),
        ("▼ 마법석",       C["section_magic"],  3),
        ("▼ 보석",         C["section_gem"],    4),
        ("▼ 마법 결정",    C["section_crystal"],12),
    ]

    row = 2
    idx = 0
    for sec_label, sec_bg, sec_count in SECTIONS:
        section_label(ws, row, len(ITEM_HEADERS), sec_label, sec_bg)
        row += 1
        for _ in range(sec_count):
            item = ITEMS[idx]; idx += 1
            values = list(item[:9])
            bg = item[9]
            for col, val in enumerate(values, 1):
                c = ws.cell(row=row, column=col, value=val)
                c.fill      = fill(bg)
                c.font      = cfont()
                c.alignment = center() if col in (4,5,6,7) else left()
                c.border    = border()
            dv_type.sqref += f"B{row}"
            dv_rt.sqref   += f"C{row}"
            ws.row_dimensions[row].height = 18
            row += 1

    ws.auto_filter.ref = f"A1:{get_column_letter(len(ITEM_HEADERS))}1"
    ws.sheet_view.showGridLines = False


# ══════════════════════════════════════════════════════════════════════════════
# 시트 2 : Recipes
# ══════════════════════════════════════════════════════════════════════════════
RECIPE_HEADERS = [
    ("recipeName",         22),
    ("resultItemName",     20),
    ("resultAmount",        9),
    ("ingredient1",        18),
    ("amount1",             7),
    ("ingredient2",        18),
    ("amount2",             7),
    ("ingredient3",        18),
    ("amount3",             7),
    ("craftTime",           9),
    ("requiredSkillType",  16),
    ("requiredSkillLevel", 14),
    ("skillGrowthAmount",  14),
    ("requiredResearchId", 20),
]

CRAFT_SKILLS = ["Labor","WeaponCrafting","ArmorCrafting","Cooking",
                "Alchemy","Research","Smelting","Stonecutting","Woodworking","Gemcutting"]

def R(name,result,ramt,i1,a1,i2,a2,i3,a3,t,skill,lv,grow,res,bg):
    return (name,result,ramt,i1,a1,i2,a2,i3,a3,t,skill,lv,grow,res,bg)

RECIPES = [
    # ── 제련 ──────────────────────────────────────────────────
    R("청동 제련", "청동 주괴",1,"구리 광석",2,"주석 광석",1,"", 0, 8.0,"Smelting", 0,0.3,"",C["section_smelt"]),
    R("강철 제련", "강철 주괴",1,"철 광석",  2,"석탄",     1,"", 0,12.0,"Smelting",20,0.5,"",C["section_smelt"]),
    R("은 제련",   "은 주괴",  1,"은 광석",  2,"",          0,"", 0,10.0,"Smelting",15,0.4,"",C["section_smelt"]),
    R("금 제련",   "금 주괴",  1,"금 광석",  2,"",          0,"", 0,14.0,"Smelting",30,0.6,"",C["section_smelt"]),
    # ── 석재 가공 ─────────────────────────────────────────────
    R("석재 가공", "다듬은 돌",2,"돌",       3,"",          0,"", 0, 5.0,"Stonecutting",0,0.3,"",C["section_stone_r"]),
    # ── 제재 ──────────────────────────────────────────────────
    R("제재",      "목재 판자",3,"목재",     2,"",          0,"", 0, 6.0,"Woodworking",0,0.3,"",C["section_wood_r"]),
    # ── 보석 가공 (원소) ──────────────────────────────────────
    R("원소 화염 결정 제작","원소 화염 결정",1,"루비",    1,"마력석", 1,"",0,10.0,"Gemcutting", 0,0.5,"",C["section_gem_r"]),
    R("원소 냉기 결정 제작","원소 냉기 결정",1,"사파이어",1,"마력석", 1,"",0,10.0,"Gemcutting", 0,0.5,"",C["section_gem_r"]),
    R("원소 번개 결정 제작","원소 번개 결정",1,"토파즈",  1,"마력석", 1,"",0,10.0,"Gemcutting", 0,0.5,"",C["section_gem_r"]),
    R("원소 독 결정 제작",  "원소 독 결정",  1,"에메랄드",1,"마력석", 1,"",0,10.0,"Gemcutting", 0,0.5,"",C["section_gem_r"]),
    # ── 보석 가공 (신성) ──────────────────────────────────────
    R("신성 화염 결정 제작","신성 화염 결정",1,"루비",    1,"에테르석",1,"",0,12.0,"Gemcutting",20,0.6,"",C["section_gem_r"]),
    R("신성 냉기 결정 제작","신성 냉기 결정",1,"사파이어",1,"에테르석",1,"",0,12.0,"Gemcutting",20,0.6,"",C["section_gem_r"]),
    R("신성 번개 결정 제작","신성 번개 결정",1,"토파즈",  1,"에테르석",1,"",0,12.0,"Gemcutting",20,0.6,"",C["section_gem_r"]),
    R("신성 독 결정 제작",  "신성 독 결정",  1,"에메랄드",1,"에테르석",1,"",0,12.0,"Gemcutting",20,0.6,"",C["section_gem_r"]),
    # ── 보석 가공 (암흑) ──────────────────────────────────────
    R("암흑 화염 결정 제작","암흑 화염 결정",1,"루비",    1,"공허석",1,"",0,12.0,"Gemcutting",20,0.6,"",C["section_gem_r"]),
    R("암흑 냉기 결정 제작","암흑 냉기 결정",1,"사파이어",1,"공허석",1,"",0,12.0,"Gemcutting",20,0.6,"",C["section_gem_r"]),
    R("암흑 번개 결정 제작","암흑 번개 결정",1,"토파즈",  1,"공허석",1,"",0,12.0,"Gemcutting",20,0.6,"",C["section_gem_r"]),
    R("암흑 독 결정 제작",  "암흑 독 결정",  1,"에메랄드",1,"공허석",1,"",0,12.0,"Gemcutting",20,0.6,"",C["section_gem_r"]),
]

def make_recipes_sheet(wb):
    ws = wb.create_sheet("Recipes")
    ws.freeze_panes = "A2"

    apply_header(ws, RECIPE_HEADERS)

    dv_skill = DataValidation(
        type="list",
        formula1='"' + ",".join(CRAFT_SKILLS) + '"',
        showDropDown=False,
    )
    ws.add_data_validation(dv_skill)

    SECTIONS_R = [
        ("▼ 제련 (용광로)",      C["section_smelt"],   4),
        ("▼ 석재 가공 (석재 가공대)", C["section_stone_r"],1),
        ("▼ 목공 (제재소)",       C["section_wood_r"],  1),
        ("▼ 보석 가공 — 원소",   C["section_gem_r"],   4),
        ("▼ 보석 가공 — 신성",   C["section_gem_r"],   4),
        ("▼ 보석 가공 — 암흑",   C["section_gem_r"],   4),
    ]

    row = 2
    idx = 0
    CENTER_COLS = {3,5,7,9,10,11,12,13}
    for sec_label, sec_bg, sec_count in SECTIONS_R:
        section_label(ws, row, len(RECIPE_HEADERS), sec_label, sec_bg)
        row += 1
        for _ in range(sec_count):
            r = RECIPES[idx]; idx += 1
            values = list(r[:14])
            bg = r[14]
            for col, val in enumerate(values, 1):
                c = ws.cell(row=row, column=col, value=val if val != 0 else "")
                c.fill      = fill(bg)
                c.font      = cfont()
                c.alignment = center() if col in CENTER_COLS else left()
                c.border    = border()
            dv_skill.sqref += f"K{row}"
            ws.row_dimensions[row].height = 18
            row += 1

    ws.auto_filter.ref = f"A1:{get_column_letter(len(RECIPE_HEADERS))}1"
    ws.sheet_view.showGridLines = False


# ══════════════════════════════════════════════════════════════════════════════
# 시트 3 : Enums (읽기 전용 레퍼런스)
# ══════════════════════════════════════════════════════════════════════════════
def make_enums_sheet(wb):
    ws = wb.create_sheet("Enums (참고)")
    ws.sheet_view.showGridLines = False

    groups = [
        ("ResourceType", RESOURCE_TYPES),
        ("CraftSkill",   CRAFT_SKILLS),
        ("WeaponType",   ["Unarmed","Shield","OneHandSword","OneHandAxe","OneHandMace",
                          "TwoHandSword","TwoHandAxe","TwoHandMace","Polearm",
                          "Bow","Crossbow","Javelin",
                          "ElementalFire","ElementalIce","ElementalLightning","ElementalPoison",
                          "DarkFire","DarkIce","DarkLightning","DarkPoison",
                          "HolyFire","HolyIce","HolyLightning","HolyPoison"]),
        ("ArmorSlot",    ["Head","Chest","Hands","Legs"]),
        ("DamageCategory",["Physical","Elemental","Dark","Holy"]),
        ("PhysicalType", ["Slash","Pierce","Blunt"]),
        ("DamageElement",["None","Fire","Ice","Lightning","Poison"]),
        ("JobCategory",  ["None","Mining","Hauling","Farming","Logging",
                          "Construction","Guard","Crafting","Research"]),
    ]

    col = 1
    for group_name, values in groups:
        ws.column_dimensions[get_column_letter(col)].width = max(len(v) for v in values) + 4
        c = ws.cell(row=1, column=col, value=group_name)
        c.fill      = fill(C["header_bg"])
        c.font      = hfont()
        c.alignment = center()
        c.border    = border()
        for r, v in enumerate(values, 2):
            c = ws.cell(row=r, column=col, value=v)
            c.fill      = fill(C["stripe"] if r % 2 == 0 else C["white"])
            c.font      = cfont()
            c.alignment = left()
            c.border    = border()
        col += 1


# ══════════════════════════════════════════════════════════════════════════════
# 시트 4 : Buildings (건물 데이터)
# ══════════════════════════════════════════════════════════════════════════════
BUILDING_HEADERS = [
    ("BuildingName",  20),
    ("Category",      16),
    ("CostType",      16),
    ("CostAmount",    10),
    ("UsesBlueprint", 13),
    ("TableType",     16),
    ("RequiresFuel",  12),
    ("FuelBurnRate",  12),
    ("FuelCapacity",  12),
    ("FuelItems",     30),
    ("메모",          30),
]

BUILD_CATEGORIES = ["Resource","Defense","Building","Misc"]

BUILDINGS = [
    # (name, category, costType, costAmt, usesBlueprint, tableType, requiresFuel, burnRate, capacity, fuelItems, note, bg)
    ("용광로",      "Building","Coal",   5, True, "Smelting",     True,  2.0,100,"목재(10), 석탄(25)","연료 필요",      C["section_smelt"]),
    ("석재 가공대", "Building","Stone",  8, True, "Stonecutting", False, 0,  0,  "",                  "",               C["section_stone_r"]),
    ("제재소",      "Building","Wood",   6, True, "Woodworking",  False, 0,  0,  "",                  "",               C["section_wood_r"]),
    ("보석 연마대", "Building","Stone",  4, True, "Gemcutting",   False, 0,  0,  "",                  "",               C["section_gem_r"]),
    ("무기 제작대", "Building","Bronze", 3, True, "WeaponCrafting",False,0,  0,  "",                  "",               C["section_weapon"]),
    ("방어구 제작대","Building","Bronze",3, True, "ArmorCrafting",False, 0,  0,  "",                  "",               C["section_armor"]),
    ("창고",        "Resource","Wood",   5, False,"",             False, 0,  0,  "",                  "즉시 배치",      C["section_stone"]),
    ("서브 코어",   "Building","MagicStone",3,True,"",            False, 0,  0,  "",                  "범위 확장+캡+1", C["section_magic"]),
]

def make_buildings_sheet(wb):
    ws = wb.create_sheet("Buildings")
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False

    apply_header(ws, BUILDING_HEADERS)

    dv_cat = DataValidation(type="list",
                            formula1='"' + ",".join(BUILD_CATEGORIES) + '"',
                            showDropDown=False)
    dv_table = DataValidation(type="list",
                              formula1='"' + ",".join(CRAFT_SKILLS + [""]) + '"',
                              showDropDown=False)
    ws.add_data_validation(dv_cat)
    ws.add_data_validation(dv_table)

    CENTER_COLS = {4,5,7,8,9}
    for row_i, b in enumerate(BUILDINGS, 2):
        values = list(b[:11])
        bg = b[11]
        for col, val in enumerate(values, 1):
            c = ws.cell(row=row_i, column=col, value=val if val != 0 else "")
            c.fill      = fill(bg)
            c.font      = cfont()
            c.alignment = center() if col in CENTER_COLS else left()
            c.border    = border()
        dv_cat.sqref   += f"B{row_i}"
        dv_table.sqref += f"F{row_i}"
        ws.row_dimensions[row_i].height = 18

    ws.auto_filter.ref = f"A1:{get_column_letter(len(BUILDING_HEADERS))}1"


# ══════════════════════════════════════════════════════════════════════════════
# 메인
# ══════════════════════════════════════════════════════════════════════════════
def main():
    wb = Workbook()
    make_items_sheet(wb)
    make_recipes_sheet(wb)
    make_buildings_sheet(wb)
    make_enums_sheet(wb)

    os.makedirs(os.path.dirname(OUTPUT_PATH), exist_ok=True)
    wb.save(OUTPUT_PATH)
    print(f"OK: {OUTPUT_PATH}")

if __name__ == "__main__":
    main()
