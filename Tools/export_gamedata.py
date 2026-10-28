# -*- coding: utf-8 -*-
"""
게임 데이터 현황표 — **에셋과 코드에서 직접 읽어** 뽑는다.

실행: python Tools/export_gamedata.py   (어디서 돌려도 된다)
출력: Tools/GameData.xlsx  (Items / Recipes / Buildings / Enums, 시트마다 자동 필터)
      Tools/Items.csv, Tools/Recipes.csv  (git에서 값 변화를 diff로 보려고)

예전 `make_gamedata.py`는 값을 **손으로 적어 둔** 기획 문서였다. 그래서 아이템이 40종에서
멈춰 있었고(실제 89종), 2026-08-02에 삭제된 `ResourceType` enum을 계속 쓰고 있었으며,
무게·스택이 에셋과 어긋나 있었다 — **표를 믿고 밸런스를 잡으면 없는 값을 잡는 셈**이다.
이제 표는 저작물이 아니라 **에셋의 그림자**다. 값을 바꿨으면 이걸 다시 돌려 표를 맞춘다.
"""
import csv, io, glob, json, os, re

os.chdir(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule

OUT = "Tools/GameData.xlsx"

# ── GUID -> 에셋 경로 ─────────────────────────────────────────────────────
guid2path = {}
for meta in glob.glob("Assets/**/*.meta", recursive=True):
    try:
        s = io.open(meta, encoding="utf-8", errors="ignore").read()
    except Exception:
        continue
    m = re.search(r"guid: (\w+)", s)
    if m:
        guid2path[m.group(1)] = meta[:-5]


def aname(p):
    return os.path.splitext(os.path.basename(p))[0]


def num(s, key, default=0.0):
    m = re.search(r"\n  " + key + r": ([-\d.]+)", s)
    return float(m.group(1)) if m else default


_ESC = {"n": "\n", "r": "\r", "t": "\t", "b": "\b", "f": "\f",
        '"': '"', "\\": "\\", "/": "/"}


def _unescape(v):
    """`\\uXXXX` · `\\n` 같은 YAML 이스케이프를 푼다.

    예전에는 `v.encode().decode("unicode_escape")` 한 줄이었다. 값이 **전부 ASCII로
    이스케이프돼 있을 때만** 맞는 방법이라(한글은 latin-1로 되읽혀 깨진다) 지금까지는
    운이 좋았을 뿐이다 — 설명처럼 긴 글에는 안 깨진 한글과 escape가 섞이기 쉽다.
    """
    out, i, n = [], 0, len(v)
    while i < n:
        if v[i] != "\\" or i + 1 >= n:
            out.append(v[i]); i += 1; continue
        nxt = v[i + 1]
        if nxt == "u" and i + 6 <= n:
            try:
                out.append(chr(int(v[i + 2:i + 6], 16))); i += 6; continue
            except ValueError:
                pass
        out.append(_ESC.get(nxt, nxt)); i += 2
    return "".join(out)


def _closed(v):
    """큰따옴표 스칼라가 이 줄 안에서 닫혔나(이스케이프된 따옴표는 안 센다)."""
    if not v.startswith('"'):
        return True
    i = 1
    while i < len(v):
        if v[i] == "\\":
            i += 2; continue
        if v[i] == '"':
            return True
        i += 1
    return False


def text(s, key):
    """`  Key: value` 한 칸.

    **Unity는 긴 문자열을 여러 줄로 접어 쓴다** — 이어 붙이지 않으면 설명이 **말 중간에서
    잘린다**(`빗아 빵을 굽는 곡물. 그대로는` 에서 끝났다). YAML은 접힌 줄바꿈을 공백
    하나로 펴므로 그렇게 잇는다. 끝을 아는 방법은 **따옴표가 닫혔는가**뿐이다 — 이어지는
    줄을 들여쓰기 깊이로 세면 `Loot:` 같은 목록의 자식 줄까지 삼킨다.
    """
    m = re.search(r"^  " + re.escape(key) + r": (.*)$", s, re.M)
    if not m:
        return ""

    v = m.group(1).rstrip()
    if not _closed(v):
        for line in s[m.end():].split("\n")[1:]:
            v += " " + line.strip()
            if _closed(v):
                break

    v = v.strip()
    if len(v) >= 2 and v[0] == v[-1] == '"':
        v = _unescape(v[1:-1])
    elif len(v) >= 2 and v[0] == v[-1] == "'":
        v = v[1:-1].replace("''", "'")

    # CSV는 git diff로 보는 것이라 한 줄로 편다 — 셀 안 줄바꿈은 diff를 읽기 어렵게 한다.
    return " ".join(v.split())


def ref(s, key):
    """`Key: {fileID: n, guid: g, type: t}` -> 에셋 이름"""
    m = re.search(r"\n  " + key + r": \{fileID: \d+, guid: (\w+)", s)
    return aname(guid2path[m.group(1)]) if m and m.group(1) in guid2path else ""


GROUP = {0: "석재", 1: "원석", 2: "보석", 3: "주괴", 4: "목재", 5: "화폐",
         6: "뼈", 7: "가죽", 8: "천"}
SLOT = {0: "머리", 1: "몸통", 2: "손", 3: "다리"}
SLOT_JSON = {"Head": 0, "Chest": 1, "Hands": 2, "Legs": 3}
WTYPE = {0: "맨손", 1: "방패", 2: "한손 검", 3: "한손 도끼", 4: "한손 둔기",
         5: "양손 검", 6: "양손 도끼", 7: "양손 둔기", 8: "장병기",
         9: "활", 10: "쇠뇌", 11: "투창"}
WTYPE_JSON = {"Unarmed": "맨손", "Shield": "방패", "OneHandSword": "한손 검",
              "OneHandAxe": "한손 도끼", "OneHandMace": "한손 둔기",
              "TwoHandSword": "양손 검", "TwoHandAxe": "양손 도끼",
              "TwoHandMace": "양손 둔기", "Polearm": "장병기",
              "Bow": "활", "Crossbow": "쇠뇌", "Javelin": "투창"}
SKILL = {0: "노동", 1: "무기제작", 2: "방어구제작", 3: "요리", 4: "연금술",
         5: "연구", 6: "제련", 7: "석공", 8: "목공", 9: "보석세공"}
BCAT = {0: "건물", 1: "생산", 2: "보관", 3: "가구", 4: "방어", 5: "바닥"}

# ── 아이템 ────────────────────────────────────────────────────────────────
items = {}

for f in glob.glob("Assets/**/*.asset", recursive=True):
    try:
        s = io.open(f, encoding="utf-8").read()
    except Exception:
        continue
    if "BaseValue:" not in s:
        continue
    cls = re.search(r"m_EditorClassIdentifier: Assembly-CSharp::(\w+)", s)
    cls = cls.group(1) if cls else "ItemData"
    name = re.search(r"m_Name: (.*)", s).group(1).strip()

    # **식별자는 `Id`가 이긴다**(`DataId.Resolve`) — 비어 있을 때만 파일명이다.
    # 안 맞추면 표를 보고 JSON 참조를 적은 사람이 조용히 못 찾는 Id를 쓰게 된다.
    explicit = re.search(r"^  Id: (\S.*)$", s, re.M)
    if explicit:
        name = explicit.group(1).strip()

    if cls == "WeaponData":
        kind, resale = WTYPE.get(int(num(s, "Type")), "지팡이"), 0.55
    elif cls == "ArmorData":
        kind, resale = SLOT.get(int(num(s, "Slot")), "?") + " 방어구", 0.5
    elif cls == "ResourceItemData":
        kind, resale = GROUP.get(int(num(s, "Group")), "자원"), 1.0
    elif cls == "FoodData":
        kind, resale = "음식", 1.0
    elif cls == "SeedData":
        kind, resale = "씨앗", 1.0
    else:
        kind, resale = "잡화", 1.0

    items[name] = dict(id=name, name=text(s, "ItemName") or name, kind=kind,
                       desc=text(s, "Description"),
                       value=int(num(s, "BaseValue")), weight=num(s, "Weight"),
                       stack=int(num(s, "Stack", 1)), w=int(num(s, "Width", 1)),
                       h=int(num(s, "Height", 1)), resale=resale, src="SO")

for f in glob.glob("Assets/StreamingAssets/Data/items/*.json"):
    d = json.load(io.open(f, encoding="utf-8"))
    t = d.get("type", "")
    if t == "weapon":
        kind, resale = WTYPE_JSON.get(d.get("weaponType", ""), "지팡이"), 0.55
    elif t == "armor":
        kind = SLOT.get(SLOT_JSON.get(d.get("armorSlot", ""), 9), "?") + " 방어구"
        resale = 0.5
    elif t == "resource":
        kind = {"Stone": "석재", "Ore": "원석", "Gem": "보석", "Ingot": "주괴",
                "Wood": "목재", "Currency": "화폐", "Bone": "뼈", "Hide": "가죽",
                "Cloth": "천"}.get(d.get("resourceGroup", ""), "자원")
        resale = 1.0
    elif t == "food":
        kind, resale = "음식", 1.0
    else:
        kind, resale = "잡화", 1.0

    items[d["id"]] = dict(id=d["id"], name=d.get("itemName", d["id"]), kind=kind,
                          desc=" ".join((d.get("description") or "").split()),
                          value=d.get("baseValue") or 0, weight=d.get("weight", 0),
                          stack=d.get("stack", 1), w=d.get("width", 1),
                          h=d.get("height", 1), resale=resale, src="JSON")

# ── 레시피 ────────────────────────────────────────────────────────────────
recipes = []
for f in glob.glob("Assets/StreamingAssets/Data/recipes/*.json"):
    d = json.load(io.open(f, encoding="utf-8"))
    recipes.append(dict(id=d.get("id", ""), name=d.get("recipeName", ""),
                        result=d.get("resultItemId"), amt=d.get("resultAmount", 1),
                        ings=[(i["itemId"], i["amount"]) for i in d.get("ingredients", [])],
                        skill=d.get("requiredSkillType", ""),
                        lvl=float(d.get("requiredSkillLevel") or 0),
                        time=float(d.get("craftTime") or 0),
                        research=d.get("requiredResearchId", ""), src="JSON"))

for f in glob.glob("Assets/Resources/Data/Recipes/*.asset"):
    s = io.open(f, encoding="utf-8").read()
    ings = []
    for m in re.finditer(r"- Item: \{fileID: \d+, guid: (\w+), type: 2\}\s*\n\s*Amount: (\d+)", s):
        p = guid2path.get(m.group(1))
        if p:
            ings.append((aname(p), int(m.group(2))))
    rm = re.search(r"\n  Result: \{fileID: \d+, guid: (\w+), type: 2\}", s)
    recipes.append(dict(id=aname(f), name=text(s, "RecipeName"),
                        result=aname(guid2path[rm.group(1)]) if rm and rm.group(1) in guid2path else "",
                        amt=int(num(s, "ResultAmount", 1)), ings=ings,
                        skill=SKILL.get(int(num(s, "RequiredSkillType")), ""),
                        lvl=num(s, "RequiredSkillLevel"), time=num(s, "CraftTime"),
                        research=ref(s, "RequiredResearch"), src="SO"))

made = {r["result"]: r for r in recipes if r["result"]}

# ── 건물 ──────────────────────────────────────────────────────────────────
buildings = []
for f in glob.glob("Assets/**/*.asset", recursive=True):
    try:
        s = io.open(f, encoding="utf-8").read()
    except Exception:
        continue
    if "Assembly-CSharp::BuildingData" not in s:
        continue
    sz = re.search(r"\n  Size: \{x: (\d+), y: (\d+)\}", s)
    buildings.append({
        "Id": aname(f),
        "이름": text(s, "BuildingName"),
        "분류": BCAT.get(int(num(s, "Category")), "?"),
        "칸": "%sx%s" % (sz.groups() if sz else ("1", "1")),
        "비용 아이템": ref(s, "CostItem"),
        "비용": int(num(s, "CostAmount")),
        "선행 연구": ref(s, "RequiredResearch"),
        "청사진": "O" if num(s, "UsesBlueprint") else "",
        "바닥": "O" if num(s, "IsFloor") else "",
        "타일 점유": "O" if num(s, "OccupiesTile") else "",
        "정렬": int(num(s, "SortOrder")),
    })
buildings.sort(key=lambda b: (b["분류"], b["정렬"], b["이름"]))

# ── enum ──────────────────────────────────────────────────────────────────
enums = []
for f in glob.glob("Assets/1_Script/**/*.cs", recursive=True):
    s = io.open(f, encoding="utf-8", errors="ignore").read()
    for m in re.finditer(r"public enum (\w+)\s*\{(.*?)\}", s, re.S):
        vals = [v.split("=")[0].strip() for v in
                re.sub(r"//[^\n]*", "", m.group(2)).split(",")]
        vals = [v for v in vals if v]
        enums.append({"enum": m.group(1), "값 수": len(vals),
                      "값": ", ".join(vals),
                      "파일": os.path.relpath(f).replace("\\", "/")})
enums.sort(key=lambda e: e["enum"])

# ── 표 만들기 ─────────────────────────────────────────────────────────────
# **찾는 규칙은 런타임과 같아야 한다.** `DataManager`의 사전이 `OrdinalIgnoreCase`라
# 대소문자만 다른 `ItemId`도 게임에서는 멀쩡히 풀린다(`Res_Bone` -> `res_bone`).
# 여기서만 대소문자를 가리면 그 재료가 **조용히 0원**이 되어 재료비가 낮게 잡히고,
# 그러면 **손해 레시피가 표에서 이득으로 보인다** — 표가 있는 이유가 그것인데.
_by_lower = {k.lower(): v for k, v in items.items()}
_missing  = set()


def item_of(iid):
    """ItemId -> 아이템. 못 찾은 것은 모아 뒀다 끝에 알린다(조용히 0원이 되지 않게)."""
    v = items.get(iid) or _by_lower.get(iid.lower())
    if v is None:
        _missing.add(iid)
    return v or {}


item_rows = []
for iid, it in items.items():
    r = made.get(iid)
    mat_cost, mat_txt = 0, ""
    if r:
        mat_txt = " + ".join("%s x%d" % (i, n) for i, n in r["ings"])
        for i, n in r["ings"]:
            mat_cost += item_of(i).get("value", 0) * n
        mat_cost /= max(1, r["amt"])

    sell = max(1, round(it["value"] * 0.5 * it["resale"])) if it["value"] else 0
    item_rows.append({
        "Id": it["id"], "이름": it["name"], "종류": it["kind"],
        "기준값(G)": it["value"], "판매가(G)": sell,
        "기준값/판매가": round(it["value"] / sell, 1) if sell else "",
        "되팔이 계수": it["resale"], "무게": it["weight"],
        "값/무게": round(it["value"] / it["weight"], 1) if it["weight"] else "",
        "스택": 20 if it["kind"] == "음식" else it["stack"],
        "칸": "%dx%d" % (it["w"], it["h"]),
        "재료": mat_txt, "재료비(G)": round(mat_cost) if r else "",
        "제작 배수": round(it["value"] / mat_cost, 2) if r and mat_cost else "",
        "출처": it["src"],
        # **설명은 맨 끝이다.** 유일하게 긴 칸이라 가운데 두면 숫자 칸들이 화면 밖으로 밀린다.
        "설명": it["desc"],
    })
order = {"화폐": 0, "석재": 1, "목재": 1, "뼈": 1, "원석": 2, "주괴": 3, "보석": 4,
         "가죽": 4, "천": 4, "씨앗": 5, "음식": 6, "잡화": 7}
item_rows.sort(key=lambda x: (order.get(x["종류"], 8), -x["기준값(G)"]))

recipe_rows = []
for r in sorted(recipes, key=lambda x: (x["skill"], x["lvl"])):
    cost = sum(item_of(i).get("value", 0) * n for i, n in r["ings"])
    out_val = item_of(r["result"]).get("value", 0) * r["amt"]
    recipe_rows.append({
        "Id": r["id"], "이름": r["name"],
        "결과": r["result"], "개수": r["amt"],
        "재료": " + ".join("%s x%d" % (i, n) for i, n in r["ings"]),
        "재료비(G)": cost, "산출값(G)": out_val,
        "배수": round(out_val / cost, 2) if cost else "",
        "스킬": r["skill"], "레벨": int(r["lvl"]), "시간(초)": r["time"],
        "선행 연구": r["research"], "출처": r["src"],
    })

# ── 특성 · 성장 배경 ──────────────────────────────────────────────────────
#
# **표시명은 코드에서 읽는다.** 에셋은 `StatName`·`JobCategory`를 **정수로** 직렬화하므로
# 여기서 한글 이름을 손으로 적으면, 스탯을 하나 더하는 순간 표가 조용히 어긋난다
# (`ResourceType` enum을 지운 이유와 같다 — 이름이 두 곳에 있으면 반드시 갈린다).


def enum_names(enum, *paths):
    """enum 선언에서 **순서**를, `=> "한글"` switch에서 **이름**을 읽어 index -> 한글.

    **둘이 다른 파일에 있을 수 있다** — `JobCategory`는 선언이 `JobCategory.cs`, 표시명은
    `JobCategoryInfo.cs`다. 한 파일만 보면 조용히 영어 이름이 나온다(실제로 그랬다).
    """
    order, ko = [], {}

    for path in paths:
        try:
            s = io.open(path, encoding="utf-8", errors="ignore").read()
        except OSError:
            continue

        m = re.search(r"public enum %s\s*\{(.*?)\}" % enum, s, re.S)
        if m and not order:
            body  = re.sub(r"//.*", "", m.group(1))
            order = [v.strip() for v in body.split(",") if v.strip()]

        ko.update(re.findall(r'%s\.(\w+)\s*=>\s*"([^"]+)"' % enum, s))

    return {i: ko.get(n, n) for i, n in enumerate(order)}


STAT_KO = enum_names("StatName", "Assets/1_Script/unit/StatGrade.cs")
JOB_KO  = enum_names("JobCategory", "Assets/1_Script/Job/JobCategory.cs",
                                    "Assets/1_Script/Job/JobCategoryInfo.cs")


def offsets(s):
    """`StatOffsets:` 블록 -> `근력 +8 · 사교 -4`"""
    m = re.search(r"\n  StatOffsets:\n((?:\s+[-\w].*\n)*)", s)
    if not m:
        return ""

    parts = []
    for stat, delta in re.findall(r"-\s*Stat:\s*(\d+)\s*\n\s*Delta:\s*([-\d.]+)", m.group(1)):
        d = float(delta)
        parts.append("%s %s%g" % (STAT_KO.get(int(stat), "?" + stat), "+" if d > 0 else "", d))
    return " · ".join(parts)


def bans(s):
    """`DisabledJobs: 0200000006000000` -> `운반 · 경비`

    **유니티는 enum 리스트를 16진 덩어리로 쓴다** — 네 바이트씩 리틀엔디언이다.
    (구조체 리스트인 `StatOffsets`는 평범한 YAML 목록인데, 원시 타입 리스트만 이렇게 나온다.)
    """
    m = re.search(r"\n  DisabledJobs: ([0-9a-fA-F]+)", s)
    if not m:
        return ""

    blob = m.group(1)
    out  = []
    for i in range(0, len(blob) - 7, 8):
        v = int.from_bytes(bytes.fromhex(blob[i:i + 8]), "little")
        out.append(JOB_KO.get(v, str(v)))
    return " · ".join(out)


traits = []
for f in sorted(glob.glob("Assets/Resources/Data/Traits/*.asset")):
    s = io.open(f, encoding="utf-8", errors="ignore").read()
    traits.append({
        "이름": text(s, "TraitName") or aname(f),
        "배타 묶음": text(s, "ExclusiveGroup"),
        "가중치": num(s, "Weight", 1.0),
        "스탯 보정": offsets(s),
        "설명": text(s, "Description"),
        "파일": aname(f),
    })

SLOT_KO = {0: "유년기", 1: "성년기"}

stories = []
for f in sorted(glob.glob("Assets/Resources/Data/Backstories/*.asset")):
    s = io.open(f, encoding="utf-8", errors="ignore").read()
    stories.append({
        "칸": SLOT_KO.get(int(num(s, "Slot", 0)), "?"),
        "묶음": text(s, "Group") or "civilian",
        "이름": text(s, "Title") or aname(f),
        "가중치": num(s, "Weight", 1.0),
        "스탯 보정": offsets(s),
        "무능력": bans(s),
        "설명": text(s, "Description"),
        "파일": aname(f),
    })

# 유년기 -> 성년기, 그 안에서 묶음별로. 표를 여는 이유가 "이 무리에 뭐가 있나"라서.
stories.sort(key=lambda r: (r["칸"] != "유년기", r["묶음"], r["이름"]))

# ── 쓰기 ──────────────────────────────────────────────────────────────────
WIDTH = {"Id": 26, "이름": 18, "종류": 13, "재료": 46, "결과": 22, "값": 90,
         "파일": 40, "enum": 20, "비용 아이템": 14, "선행 연구": 18, "칸": 6, "출처": 7,
         "설명": 64, "스탯 보정": 34, "배타 묶음": 12, "묶음": 11, "무능력": 14}
NUMERIC = {"기준값(G)", "판매가(G)", "기준값/판매가", "되팔이 계수", "무게", "값/무게",
           "스택", "재료비(G)", "제작 배수", "개수", "산출값(G)", "배수", "레벨",
           "시간(초)", "비용", "정렬", "값 수", "가중치"}

wb = Workbook()
wb.remove(wb.active)


def sheet(title, rows, freeze="C2"):
    ws = wb.create_sheet(title)
    cols = list(rows[0].keys())
    ws.append(cols)
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2D3D")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = WIDTH.get(c, 12)

    for r in rows:
        out = []
        for c in cols:
            v = r[c]
            # 숫자는 숫자로 넣는다 — 문자로 넣으면 정렬·필터가 사전순이 된다.
            if c in NUMERIC and v not in ("", None):
                try:
                    v = float(v)
                    if v == int(v):
                        v = int(v)
                except (TypeError, ValueError):
                    pass
            out.append(v)
        ws.append(out)

    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(cols)), ws.max_row)
    ws.freeze_panes = freeze
    return ws, cols


ws, cols = sheet("Items", item_rows)


def col(cols, name):
    return get_column_letter(cols.index(name) + 1)


# 제작 배수 1 미만 = **만들수록 손해**. 표에서 가장 먼저 눈에 띄어야 하는 값이다.
for c in ("제작 배수",):
    L = col(cols, c)
    ws.conditional_formatting.add(
        "%s2:%s%d" % (L, L, ws.max_row),
        CellIsRule(operator="lessThan", formula=["1"],
                   fill=PatternFill("solid", fgColor="FFC7CE"),
                   font=Font(color="9C0006", bold=True)))
L = col(cols, "기준값(G)")
ws.conditional_formatting.add(
    "%s2:%s%d" % (L, L, ws.max_row),
    ColorScaleRule(start_type="min", start_color="FFFFFF",
                   end_type="max", end_color="63BE7B"))

ws2, cols2 = sheet("Recipes", recipe_rows)
L = col(cols2, "배수")
ws2.conditional_formatting.add(
    "%s2:%s%d" % (L, L, ws2.max_row),
    CellIsRule(operator="lessThan", formula=["1"],
               fill=PatternFill("solid", fgColor="FFC7CE"),
               font=Font(color="9C0006", bold=True)))

sheet("Buildings", buildings)

# **비어 있으면 시트를 안 만든다** — `Backstories`는 에디터 메뉴를 한 번 돌려야 생긴다
# (`DungeonColony → Create Backstories`). 빈 시트를 남기면 "에셋이 없다"가 아니라
# "표가 고장났다"로 보인다.
if traits:  sheet("Traits", traits, freeze="B2")
if stories: sheet("Backstories", stories, freeze="C2")

sheet("Enums", enums, freeze="B2")

wb.save(OUT)

for name, rows in (("Items", item_rows), ("Recipes", recipe_rows),
                   ("Traits", traits), ("Backstories", stories)):
    if not rows:
        continue
    with io.open("Tools/%s.csv" % name, "w", encoding="utf-8-sig", newline="") as fp:
        w = csv.DictWriter(fp, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)

print("%s : items %d / recipes %d / buildings %d / traits %d / backstories %d / enums %d"
      % (OUT, len(item_rows), len(recipe_rows), len(buildings),
         len(traits), len(stories), len(enums)))
if not stories:
    print("  (성장 배경 에셋이 없습니다 — 유니티에서 DungeonColony > Create Backstories 를 한 번 도세요)")
# **못 찾은 ItemId는 반드시 알린다.** 오타는 조용히 아무것도 안 나오는 것으로 나타나고
# (레시피는 재료가 0원이 되고 게임에서는 재료가 없는 것이 된다) 표만 보면 티가 안 난다.
if _missing:
    print("  못 찾은 ItemId %d개: %s" % (len(_missing), ", ".join(sorted(_missing))))
